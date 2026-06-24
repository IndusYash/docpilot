from docx.shared import Pt
from docx.enum.text import (
    WD_PARAGRAPH_ALIGNMENT
)


class StyleEngine:

    @staticmethod
    def apply_title(
        document,
        text
    ):

        paragraph = (
            document.add_paragraph()
        )

        paragraph.alignment = (
            WD_PARAGRAPH_ALIGNMENT.CENTER
        )

        run = paragraph.add_run(
            text
        )

        run.bold = True

        run.font.name = (
            "Times New Roman"
        )

        run.font.size = Pt(20)

    @staticmethod
    def apply_heading(
        document,
        text
    ):

        paragraph = (
            document.add_paragraph()
        )

        paragraph.alignment = (
            WD_PARAGRAPH_ALIGNMENT.LEFT
        )

        run = paragraph.add_run(
            text
        )

        run.bold = True

        run.font.name = (
            "Times New Roman"
        )

        run.font.size = Pt(14)

    @staticmethod
    def apply_subheading(
        document,
        text
    ):

        paragraph = (
            document.add_paragraph()
        )

        paragraph.alignment = (
            WD_PARAGRAPH_ALIGNMENT.LEFT
        )

        run = paragraph.add_run(
            text
        )

        run.bold = True

        run.font.name = (
            "Times New Roman"
        )

        run.font.size = Pt(12)

    @staticmethod
    def apply_body_text(
        document,
        text
    ):

        paragraph = (
            document.add_paragraph()
        )

        paragraph.alignment = (
            WD_PARAGRAPH_ALIGNMENT.JUSTIFY
        )

        paragraph.paragraph_format.space_after = Pt(
            6
        )

        paragraph.paragraph_format.line_spacing = (
            1.15
        )

        run = paragraph.add_run(
            text
        )

        run.font.name = (
            "Times New Roman"
        )

        run.font.size = Pt(12)